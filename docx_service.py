import os

from docx import Document as DocxDocument
from docx.document import Document as _DocxDocumentType
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle


def _iter_block_items(parent):
    """Yield paragraphs and tables in document order (python-docx has no
    built-in API for this ??? .paragraphs and .tables are separate, unordered
    relative to each other)."""
    parent_elm = parent.element.body if isinstance(parent, _DocxDocumentType) else parent._tc
    for child in parent_elm.iterchildren():
        if child.tag == qn('w:p'):
            yield DocxParagraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield DocxTable(child, parent)


def _escape(text: str) -> str:
    return (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


class DocxConversionError(Exception):
    """Raised when a .docx file can't be opened/converted - invalid,
    corrupt, truncated, or not actually a Word package (e.g. a renamed
    file of a different type). Callers should catch this and surface a
    clean, user-facing message rather than the raw python-docx traceback."""


def docx_to_pdf(docx_path: str, pdf_path: str) -> None:
    """
    Render a .docx file's text content to a PDF so it can flow through the
    same PDF-based pipeline as every other upload (page rendering, text
    extraction, viewer, export, highlighting). Best-effort text fidelity ???
    paragraphs, headings and tables are preserved; images and complex layout
    are not, since the goal is accurate field extraction, not a visual clone.
    """
    try:
        src = DocxDocument(docx_path)
    except PackageNotFoundError as e:
        raise DocxConversionError(f"'{os.path.basename(docx_path)}' is not a valid .docx file: {e}") from e
    except Exception as e:
        raise DocxConversionError(f"Could not open '{os.path.basename(docx_path)}': {e}") from e

    style_body = ParagraphStyle(
        "body", fontName="Helvetica", fontSize=10.5,
        textColor=colors.black, leading=15, spaceAfter=2.5 * mm,
    )
    style_heading = ParagraphStyle(
        "heading", fontName="Helvetica-Bold", fontSize=13,
        textColor=colors.black, leading=17,
        spaceBefore=4 * mm, spaceAfter=2.5 * mm,
    )

    story = []
    for block in _iter_block_items(src):
        if isinstance(block, DocxParagraph):
            text = block.text.strip()
            if not text:
                story.append(Spacer(1, 2 * mm))
                continue
            is_heading = (block.style.name or "").lower().startswith("heading") or \
                (block.style.name or "").lower() == "title"
            story.append(Paragraph(_escape(text), style_heading if is_heading else style_body))
        elif isinstance(block, DocxTable):
            rows = [[_escape(cell.text.strip()) for cell in row.cells] for row in block.rows]
            rows = [r for r in rows if any(r)]
            if not rows:
                continue
            cell_style = ParagraphStyle("cell", fontName="Helvetica", fontSize=9, leading=12)
            table_data = [[Paragraph(c, cell_style) for c in r] for r in rows]
            t = Table(table_data, repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#b3b3b3")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(t)
            story.append(Spacer(1, 3 * mm))

    if not story:
        story = [Paragraph("(empty document)", style_body)]

    os.makedirs(os.path.dirname(pdf_path), exist_ok=True)
    doc = SimpleDocTemplate(
        pdf_path, pagesize=A4,
        leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=20 * mm, bottomMargin=20 * mm,
    )
    doc.build(story)
